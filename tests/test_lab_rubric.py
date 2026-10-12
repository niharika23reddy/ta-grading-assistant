import io
import shutil

import docx
import pytest
from PIL import Image, ImageDraw, ImageFont

from grader.checklist import Checklist, attach_instruction_context, find_conflicts, read_rubric_csv
from grader.extract import extract_file, has_unread_content
from grader.models import Decision
from grader.scoring import score

# Synthetic rubric in the same layout as a level-based LMS rubric export.
LEVEL_RUBRIC = """Criteria,Expected,Expected Pts,Satisfactory,Satisfactory Pts,Insufficient,Insufficient Pts,Max Pts
Lab Exercise 1,"1. Screenshot of the command and its output. 2. Description of the command.",2,Missing description or screenshot.,1,Missing both.,0,2
Lab Exercise 2,All tools installed with screenshots.,3,"Some tools missing; each tool is worth 0.5 points.",1.5,Nothing installed.,0,3
Total,,,,,,,5
"""


def dec(req_id, status):
    return Decision(req_id, "", status, "Correct", [], "", "", False)


def test_level_rubric_criteria_levels_and_total():
    r = read_rubric_csv(LEVEL_RUBRIC)
    c1, c2 = r.criteria
    assert (c1.description, c1.max_points, c1.levels) == ("Lab Exercise 1", 2.0, {"full": 2, "partial": 1, "none": 0})
    assert "Missing description or screenshot" in c1.notes
    assert c2.levels == {"full": 3, "partial": 1.5, "none": 0}
    assert r.total == 5


def test_level_rubric_drafts_linked_requirements_from_expected_column():
    r = read_rubric_csv(LEVEL_RUBRIC)
    assert [(q.req_id, q.criterion_id, q.description) for q in r.requirements] == [
        ("R1", "C1", "Lab Exercise 1: Screenshot of the command and its output"),
        ("R2", "C1", "Lab Exercise 1: Description of the command"),
        ("R3", "C2", "Lab Exercise 2: All tools installed with screenshots")]


def test_per_item_rule_is_flagged_and_partial_stays_provisional():
    r = read_rubric_csv(LEVEL_RUBRIC)
    cl = Checklist("X", "A", r.requirements, r.criteria)
    issues = find_conflicts(cl, rubric_total=r.total)
    assert len(issues) == 1 and "per-item" in issues[0] and "C2" in issues[0]
    s = score(r.criteria, r.requirements, [dec("R1", "Met"), dec("R2", "Missing"), dec("R3", "Partially met")])
    c1, c2 = s.criteria
    assert c1.suggested_points == 1 and not c1.provisional  # one of two items missing = Satisfactory
    assert c2.suggested_points == 1.5 and c2.provisional
    assert s.suggested_total == 2.5


def test_total_row_mismatch_reported():
    bad = LEVEL_RUBRIC.replace(",5\n", ",6\n")
    r = read_rubric_csv(bad)
    issues = find_conflicts(Checklist("X", "A", r.requirements, r.criteria), rubric_total=r.total)
    assert any("Total row says 6" in i for i in issues)


def test_template_rubric_still_supported():
    r = read_rubric_csv("criterion_id,description,max_points,full_points,partial_points,none_points\n"
                        "C1,Intro,2,2,,0\n")
    assert [(c.criterion_id, c.levels) for c in r.criteria] == [("C1", {"full": 2.0, "none": 0.0})]


def test_instruction_sections_become_requirement_notes():
    r = read_rubric_csv(LEVEL_RUBRIC)
    instructions = ("Week 2 lab\nLab Exercise 1\nRun `uname -a` and explain the output.\n"
                    "Lab Exercise 2\nInstall git, python and docker.\n")
    attach_instruction_context(r.requirements, r.criteria, instructions)
    assert "uname -a" in r.requirements[0].notes and "docker" not in r.requirements[0].notes
    assert "docker" in r.requirements[2].notes


def _screenshot_png(text):
    img = Image.new("RGB", (1100, 120), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSansMono.ttf", 32)
    except OSError:
        font = ImageFont.load_default(size=32)
    draw.text((10, 40), text, fill="black", font=font)
    png = io.BytesIO()
    img.save(png, format="PNG")
    png.seek(0)
    return png


def _docx(parts):
    """parts: strings become paragraphs, ("img", text) becomes a pasted screenshot."""
    d = docx.Document()
    for part in parts:
        if isinstance(part, tuple):
            d.add_picture(_screenshot_png(part[1]))
        else:
            d.add_paragraph(part)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_with_screenshot(text):
    return _docx(["Exercise 1 screenshot:", ("img", text)])


def test_word_screenshots_flagged_when_ocr_off():
    [rep] = extract_file("lab.docx", _docx_with_screenshot("uname -a Linux ubuntu"), ocr=False)
    assert has_unread_content([rep])
    assert "1 image(s)" in rep.pages[0].warnings[0]
    # the model is still told a screenshot is there
    assert "[Screenshot 1: image present; not read" in rep.pages[0].text


needs_ocr = pytest.mark.skipif(shutil.which("tesseract") is None, reason="needs the tesseract OCR binary")


@needs_ocr
def test_word_screenshots_read_with_ocr_in_place():
    data = _docx(["Lab Exercise 1", "I ran uname to see the kernel.", ("img", "$ uname -a Linux ubuntu 6.8"),
                  "Lab Exercise 2", "I checked the disk.", ("img", "$ df -h  Filesystem Size Used Avail")])
    [rep] = extract_file("lab.docx", data, ocr=True)
    assert len(rep.pages) == 1 and not has_unread_content([rep])
    text = rep.pages[0].text
    # each screenshot's OCR text sits right after the exercise it belongs to
    assert text.index("Lab Exercise 1") < text.index("[Screenshot 1: OCR text") < text.index("uname -a") \
        < text.index("Lab Exercise 2") < text.index("[Screenshot 2: OCR text") < text.index("df -h")
    assert any("2 screenshot(s) read with OCR" in w for w in rep.pages[0].warnings)


@needs_ocr
def test_screenshot_chunks_are_marked_as_ocr():
    from grader.chunking import chunk_reports

    [rep] = extract_file("lab.docx", _docx(["Lab Exercise 1", ("img", "$ aescrypt -h usage aescrypt")]), ocr=True)
    [chunk] = chunk_reports([rep], "p")
    assert chunk.ocr and "(OCR)" in chunk.location()


class _Scripted:
    name = "scripted"

    def __init__(self, answer):
        self.answer, self.prompts = answer, []

    def judge(self, system, user, requirement, chunks):
        self.prompts.append((system, user))
        return dict(self.answer)


MISSING = {"status": "Missing", "correctness": "Not assessed", "evidence": [], "reasoning": "no screenshot"}


@needs_ocr
def test_exercise_screenshot_reaches_the_model(store, settings):
    """The case from real grading: the screenshot of 'aescrypt -h' is pasted under Exercise 1."""
    from grader.evaluate import SYSTEM_PROMPT, evaluate_submission
    from grader.models import Requirement, Scope
    from grader.pipeline import ingest_submission

    filler = [f"Note {i}: notes about virtual machines, networking and snapshots." for i in range(120)]
    data = _docx(filler + ["Lab Exercise 1", "The command prints the help text for the AES tool.",
                           ("img", "$ aescrypt -h  usage: aescrypt -e|-d")] + filler)
    scope = Scope("CS", "W2", "S001", "v1")
    ingest_submission(store, scope, [("lab.docx", data)], settings)
    req = Requirement("R1", "Lab Exercise 1: Screenshot with the given command and expected output",
                      "C1", notes="Instructions: Run aescrypt -h in the terminal and take a screenshot.")
    p = _Scripted(MISSING)
    [d], _ = evaluate_submission([req], scope, store, p, unread_content=False, top_k=2)
    system, user = p.prompts[0]
    assert "PROVES a" in system and "contains 1 screenshot(s)" in user
    assert "aescrypt -h" in user and "contains screenshot text" in user
    # a "Missing" verdict on a screenshot requirement always tells the TA to look at the screenshots
    assert d.ta_review and any("confirm none of them shows" in r for r in d.review_reasons)


def test_missing_screenshot_unverified_when_no_screenshot_was_checked(monkeypatch):
    from grader.evaluate import evaluate_requirement
    from grader.models import Chunk, Requirement, Scope

    shot = Chunk("c2", "[Screenshot 1: OCR text of the image begins]\nzzz\n[Screenshot 1: OCR text of the image ends]",
                 "lab.docx", None, 3, 5, 0, 10, ocr=True)
    text = Chunk("c1", "Lab Exercise 1 description of the command", "lab.docx", None, 1, 2, 0, 10)

    class Store:  # vector search returns only the text chunk
        def search(self, scope, query, k):
            return [text]

        def reference_chunks(self, *a):
            return []

    req = Requirement("R1", "Lab Exercise 1: Screenshot with the given command", "C1")
    d = evaluate_requirement(req, Scope("CS", "W2", "S001", "v1"), Store(), _Scripted(MISSING),
                             all_chunks=[text, shot], unread_content=False, injection_hits=[])
    # a screenshot requirement is always shown the related screenshots, so "Missing" can stand
    assert d.status == "Missing" and d.ta_review

    req2 = Requirement("R2", "Lab Exercise 1: picture with the given command", "C1")
    d2 = evaluate_requirement(req2, Scope("CS", "W2", "S001", "v1"), Store(), _Scripted(MISSING),
                              all_chunks=[text], unread_content=False, injection_hits=[])
    assert d2.status == "Missing"  # no screenshots anywhere: absence stands, TA still reviews

    # if no screenshot is among the checked passages, "Missing" is not trusted
    monkeypatch.setattr("grader.evaluate.broad_search", lambda *a, **k: [])
    monkeypatch.setattr("grader.evaluate.screenshot_candidates", lambda *a, **k: [])
    d3 = evaluate_requirement(req, Scope("CS", "W2", "S001", "v1"), Store(), _Scripted(MISSING),
                              all_chunks=[text, shot], unread_content=False, injection_hits=[])
    assert d3.status == "Unable to verify"
    assert any("none were among the passages" in r for r in d3.review_reasons)
