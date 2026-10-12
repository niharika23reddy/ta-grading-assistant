import io
import zipfile

import docx

from grader.extract import extract_file, has_unread_content, unreadable_parts
from tests.conftest import make_pdf


def test_pdf_text_and_scanned_page_flagged():
    data = make_pdf(["My research question is about rentals and weather.", None])
    [rep] = extract_file("report.pdf", data, ocr=False)
    assert rep.readable
    assert rep.pages[0].method == "text" and "research question" in rep.pages[0].text
    assert rep.pages[1].method == "none"
    warnings = " ".join(unreadable_parts([rep]))
    assert "page 2" in warnings and "not read" in warnings
    assert has_unread_content([rep])


def test_docx_text_and_tables():
    d = docx.Document()
    d.add_paragraph("The dataset has 731 rows.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "mean", "4504"
    buf = io.BytesIO()
    d.save(buf)
    [rep] = extract_file("r.docx", buf.getvalue())
    assert "731 rows" in rep.text and "mean | 4504" in rep.text
    assert not has_unread_content([rep])


def test_unsupported_and_binary_files_flagged():
    [rep] = extract_file("video.mp4", b"\x00\x01")
    assert not rep.readable and "Unsupported" in rep.warnings[0]
    [rep] = extract_file("data.txt", b"abc\x00def")
    assert not rep.readable


def test_zip_unpacks_members_and_skips_path_traversal():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("sub/analysis.py", "def clean_data(df):\n    return df\n")
        z.writestr("../evil.txt", "should be skipped")
    reps = extract_file("submission.zip", buf.getvalue())
    assert [r.file for r in reps] == ["submission.zip/sub/analysis.py"]


def test_image_without_ocr_is_flagged_not_read():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (50, 50), "white").save(buf, format="PNG")
    [rep] = extract_file("screenshot.png", buf.getvalue(), ocr=False)
    assert not rep.readable and has_unread_content([rep])
