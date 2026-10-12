import io

import docx

from grader.checklist import Checklist, draft_requirements, draft_rubric_from_text, find_conflicts
from grader.extract import extract_file
from grader.models import Criterion, Decision, Requirement
from grader.scoring import approved_total, score


def dec(req_id, status, correctness="Correct"):
    return Decision(req_id, "", status, correctness, [], "", "", False)


CRITS = [Criterion("C1", "a", 4, {"full": 4, "partial": 2, "none": 0}),
         Criterion("C2", "b", 3, {"full": 3, "none": 0}),  # no partial rule
         Criterion("C3", "c", 3.5, {})]
REQS = [Requirement("R1", "x", "C1"), Requirement("R2", "y", "C1"),
        Requirement("R3", "z", "C2"), Requirement("R4", "w", "C3")]


def test_full_marks_sum():
    s = score(CRITS, REQS, [dec("R1", "Met"), dec("R2", "Met"), dec("R3", "Met"), dec("R4", "Met")])
    assert s.suggested_total == 10.5 and s.max_total == 10.5 and s.complete and not s.provisional


def test_partial_uses_only_rubric_rule():
    s = score(CRITS, REQS, [dec("R1", "Met"), dec("R2", "Missing"), dec("R3", "Partially met"), dec("R4", "Met")])
    c1, c2, c3 = s.criteria
    assert c1.level == "partial" and c1.suggested_points == 2
    assert c2.level == "partial" and c2.suggested_points is None and c2.provisional
    assert s.suggested_total == 2 + 3.5 and s.unresolved == ["C2"] and not s.complete


def test_unable_to_verify_is_unresolved_not_zero():
    s = score(CRITS, REQS, [dec("R1", "Met"), dec("R2", "Unable to verify"), dec("R3", "Missing"), dec("R4", "Met")])
    assert s.criteria[0].suggested_points is None and s.criteria[1].suggested_points == 0
    assert s.suggested_total == 3.5


def test_unable_to_verify_can_be_counted_as_missing():
    decisions = [dec("R1", "Met"), dec("R2", "Unable to verify"), dec("R3", "Missing"), dec("R4", "Met")]
    s = score(CRITS, REQS, decisions, unverified_as_missing=True)
    c1 = s.criteria[0]
    # R1 met, R2 counted as missing: partial credit by the rubric, flagged for the TA
    assert c1.level == "partial" and c1.suggested_points is not None and c1.provisional
    assert any("counted as missing" in n for n in c1.notes)
    assert s.complete


def test_incorrect_work_is_partial_and_not_assessed_is_provisional():
    s = score(CRITS, REQS, [dec("R1", "Met", "Incorrect"), dec("R2", "Met"), dec("R3", "Met", "Not assessed"),
                            dec("R4", "Met")])
    assert s.criteria[0].level == "partial" and s.criteria[0].suggested_points == 2
    assert s.criteria[1].suggested_points == 3 and s.criteria[1].provisional


def test_float_arithmetic_rounds_cleanly():
    crits = [Criterion(f"C{i}", "", 0.1, {}) for i in range(3)]
    reqs = [Requirement(f"R{i}", "", f"C{i}") for i in range(3)]
    s = score(crits, reqs, [dec(f"R{i}", "Met") for i in range(3)])
    assert s.suggested_total == 0.3 and s.max_total == 0.3


def test_approved_total_requires_every_criterion_and_valid_range():
    assert approved_total({"C1": 4, "C2": 3, "C3": 3.5}, CRITS) == (10.5, [])
    total, problems = approved_total({"C1": 4, "C2": None, "C3": 9}, CRITS)
    assert total is None and len(problems) == 2


def test_draft_requirements_from_sample(sample_checklist):
    reqs = sample_checklist.requirements
    assert len(reqs) == 7
    assert reqs[4].description.startswith("Implement a function named clean_data")


def test_conflicts_detected():
    cl = Checklist("X", "A1", [Requirement("R1", "x", "C1"), Requirement("R2", "y", ""),
                               Requirement("R3", "z", "C9")],
                   [Criterion("C1", "a", 5, {"full": 5, "partial": 7}), Criterion("C2", "b", 5, {})])
    issues = " ".join(find_conflicts(cl, "This assignment is worth 20 points."))
    assert "R2 is not linked" in issues and "C9" in issues and "C2 (b) has no matching" in issues
    assert "outside 0-5" in issues and "total of 20" in issues


def test_sample_checklist_is_consistent(sample_checklist):
    text = open("samples/assignment/instructions.md").read()
    issues = find_conflicts(sample_checklist, text)
    assert issues == ["C2 has no partial-credit rule; partial work will be left for the TA."]


def test_rubric_text_draft_never_invents_partial_credit():
    crits = draft_rubric_from_text("1. Introduction (2 points)\n2. Analysis - 5 pts\nTotal of 7 points")
    assert [(c.description, c.max_points, c.levels) for c in crits] == [
        ("Introduction", 2.0, {}), ("Analysis", 5.0, {})]


def test_checklist_json_roundtrip(sample_checklist):
    again = Checklist.from_json(sample_checklist.to_json())
    assert again.to_json() == sample_checklist.to_json()


def test_word_numbered_lists_and_plain_instructions_are_drafted():
    d = docx.Document()
    d.add_heading("Virtual Lab Set Up", 1)
    d.add_paragraph("In this lab you will set up your virtual machine.")
    d.add_paragraph("Download and install VirtualBox on your computer.", style="List Number")
    d.add_paragraph("Take a screenshot of the VM running.", style="List Bullet")
    d.add_paragraph("You must submit a single Word document.")
    buf = io.BytesIO()
    d.save(buf)
    [rep] = extract_file("lab.docx", buf.getvalue())
    assert [r.description for r in draft_requirements(rep.text)] == [
        "Download and install VirtualBox on your computer.",
        "Take a screenshot of the VM running.",
        "You must submit a single Word document."]


def test_rubric_table_rows_become_criteria_without_partial_credit():
    d = docx.Document()
    t = d.add_table(rows=4, cols=3)
    rows = [("Criteria", "Description", "Points"), ("VM installed", "VirtualBox runs", "5"),
            ("Screenshot", "Shows the desktop", "3 pts"), ("Total", "", "8")]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            t.cell(i, j).text = val
    buf = io.BytesIO()
    d.save(buf)
    [rep] = extract_file("rubric.docx", buf.getvalue())
    crits = draft_rubric_from_text(rep.text)
    assert [(c.description, c.max_points, c.levels) for c in crits] == [
        ("VM installed", 5.0, {}), ("Screenshot", 3.0, {})]
