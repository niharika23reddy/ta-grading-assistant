"""Text extraction with explicit quality warnings.

Nothing here executes student files. Code and notebooks are read as plain text.
Pages that cannot be read are reported, never silently dropped, so the grader
can mark related requirements "Unable to verify" instead of "Missing".
"""
from __future__ import annotations

import io
import json
import shutil
import zipfile
from pathlib import Path, PurePosixPath

from .models import ExtractionReport, PageText

TEXT_EXTS = {
    ".txt", ".md", ".py", ".java", ".c", ".cpp", ".h", ".hpp", ".js", ".ts",
    ".sql", ".r", ".csv", ".tsv", ".json", ".html", ".css", ".m", ".sh", ".yaml", ".yml",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"}
MIN_PAGE_CHARS = 25  # below this a PDF page is treated as image-only / scanned
MAX_ZIP_MEMBERS = 200
MAX_ZIP_BYTES = 100 * 1024 * 1024


def ocr_available() -> bool:
    if shutil.which("tesseract") is None:
        return False
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return True


def _ocr_image(img) -> str:
    import pytesseract

    return pytesseract.image_to_string(img)


def _ocr_page(name: str, page_no: int | None, img, ocr: bool) -> PageText:
    if not ocr:
        return PageText(name, page_no, "", "none",
                        ["Image-only content; OCR is off, so it was not read. TA must check it."])
    if not ocr_available():
        return PageText(name, page_no, "", "none",
                        ["Image-only content; Tesseract is not installed, so it was not read. TA must check it."])
    text = _ocr_image(img).strip()
    if len(text) < MIN_PAGE_CHARS:
        return PageText(name, page_no, text, "ocr",
                        ["OCR found little or no text (handwriting, diagram or poor scan). TA must check it."])
    return PageText(name, page_no, text, "ocr",
                    ["Read with OCR; wording may contain recognition errors. Verify quoted evidence."])


def extract_pdf(name: str, data: bytes, ocr: bool) -> ExtractionReport:
    import pdfplumber

    pages: list[PageText] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if len(text) >= MIN_PAGE_CHARS:
                    warn = []
                    if page.images:
                        warn.append(f"Page has {len(page.images)} embedded image(s) that were not read.")
                    pages.append(PageText(name, i, text, "text", warn))
                    continue
                img = _render_pdf_page(data, i - 1)
                pt = _ocr_page(name, i, img, ocr) if img is not None else PageText(
                    name, i, text, "none", ["Page could not be rendered for OCR."])
                if text and not pt.text:
                    pt.text = text
                pt.warnings.insert(0, "Page has little or no selectable text (likely scanned or a screenshot).")
                pages.append(pt)
    except Exception as exc:  # corrupt or encrypted PDFs
        return ExtractionReport(name, "pdf", [], False, [f"PDF could not be opened: {exc}"])
    return _report(name, "pdf", pages)


def _render_pdf_page(data: bytes, index: int):
    try:
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(data)
        return doc[index].render(scale=2).to_pil()
    except Exception:
        return None


def _docx_images(document) -> list[bytes]:
    """Embedded pictures in document order (screenshots pasted into the report)."""
    blobs = []
    for shape in document.inline_shapes:
        try:
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            blobs.append(document.part.related_parts[rid].blob)
        except Exception:
            blobs.append(b"")
    return blobs


def extract_docx(name: str, data: bytes, ocr: bool = True) -> ExtractionReport:
    import docx
    from PIL import Image

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return ExtractionReport(name, "docx", [], False, [f"DOCX could not be opened: {exc}"])
    lines = []
    for para in document.paragraphs:
        # Word's automatic numbering and bullets are not part of the paragraph text,
        # so mark list items explicitly; checklist drafting relies on them.
        style = (para.style.name if para.style is not None else "") or ""
        is_list = style.lower().startswith("list") or para._p.pPr is not None and para._p.pPr.numPr is not None
        lines.append(f"- {para.text}" if is_list and para.text.strip() else para.text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(lines).strip()
    pages = [PageText(name, None, text, "text", [])]

    images = _docx_images(document)
    if images and not (ocr and ocr_available()):
        why = "OCR is off" if not ocr else "Tesseract is not installed"
        pages[0].warnings.append(f"Document contains {len(images)} image(s) (screenshots, figures) that were not "
                                 f"read because {why}. TA must check them.")
    elif images:
        for k, blob in enumerate(images, start=1):
            label = f"{name} (image {k})"
            try:
                img = Image.open(io.BytesIO(blob))
                img.load()
            except Exception:
                pages.append(PageText(label, None, "", "none", ["Image format could not be opened. TA must check it."]))
                continue
            pages.append(_ocr_page(label, None, img, ocr))
    return _report(name, "docx", pages)


def extract_text(name: str, data: bytes) -> ExtractionReport:
    warnings = []
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="replace")
        warnings.append("File is not valid UTF-8; some characters were replaced.")
    if "\x00" in text:
        return ExtractionReport(name, "binary", [], False, ["File looks binary, not text. It was not read."])
    return _report(name, Path(name).suffix.lstrip(".").lower() or "text",
                   [PageText(name, None, text, "text", warnings)])


def extract_ipynb(name: str, data: bytes) -> ExtractionReport:
    try:
        nb = json.loads(data.decode("utf-8"))
    except Exception as exc:
        return ExtractionReport(name, "ipynb", [], False, [f"Notebook could not be parsed: {exc}"])
    parts, warnings, n_images = [], [], 0
    for idx, cell in enumerate(nb.get("cells", []), start=1):
        src = "".join(cell.get("source", []))
        parts.append(f"# [cell {idx}: {cell.get('cell_type', '?')}]\n{src}")
        for out in cell.get("outputs", []):
            if "text" in out:
                parts.append("# [output]\n" + "".join(out["text"]))
            data_out = out.get("data", {})
            if "text/plain" in data_out:
                parts.append("# [output]\n" + "".join(data_out["text/plain"]))
            if any(k.startswith("image/") for k in data_out):
                n_images += 1
    if n_images:
        warnings.append(f"Notebook has {n_images} image output(s) (plots) that were not read. TA must check them.")
    warnings.append("Notebook outputs are as saved by the student; the code was not re-run.")
    return _report(name, "ipynb", [PageText(name, None, "\n".join(parts), "text", warnings)])


def extract_image(name: str, data: bytes, ocr: bool) -> ExtractionReport:
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as exc:
        return ExtractionReport(name, "image", [], False, [f"Image could not be opened: {exc}"])
    return _report(name, "image", [_ocr_page(name, None, img, ocr)])


def extract_zip(name: str, data: bytes, ocr: bool) -> list[ExtractionReport]:
    reports = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return [ExtractionReport(name, "zip", [], False, ["ZIP archive is corrupt."])]
    members = [m for m in zf.infolist() if not m.is_dir()]
    if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_ZIP_BYTES:
        return [ExtractionReport(name, "zip", [], False, ["ZIP archive is too large to unpack safely."])]
    for m in members:
        path = PurePosixPath(m.filename)
        if path.is_absolute() or ".." in path.parts or path.name.startswith(".") or "__MACOSX" in path.parts:
            continue
        reports.extend(extract_file(f"{name}/{m.filename}", zf.read(m), ocr))
    return reports


def extract_file(name: str, data: bytes, ocr: bool = True) -> list[ExtractionReport]:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return [extract_pdf(name, data, ocr)]
    if ext == ".docx":
        return [extract_docx(name, data, ocr)]
    if ext == ".ipynb":
        return [extract_ipynb(name, data)]
    if ext in IMAGE_EXTS:
        return [extract_image(name, data, ocr)]
    if ext == ".zip":
        return extract_zip(name, data, ocr)
    if ext in TEXT_EXTS:
        return [extract_text(name, data)]
    return [ExtractionReport(name, ext.lstrip(".") or "unknown", [], False,
                             [f"Unsupported file type '{ext or 'none'}'. It was not read. TA must check it."])]


def _report(name: str, file_type: str, pages: list[PageText]) -> ExtractionReport:
    readable = any(p.text.strip() for p in pages)
    warnings = [] if readable else ["No readable text was found in this file."]
    return ExtractionReport(name, file_type, pages, readable, warnings)


def unreadable_parts(reports: list[ExtractionReport]) -> list[str]:
    """Human-readable list of anything that could not be read fully."""
    out = []
    for r in reports:
        for w in r.warnings:
            out.append(f"{r.file}: {w}")
        for p in r.pages:
            where = p.file if p.page is None else f"{p.file} page {p.page}"
            for w in p.warnings:
                out.append(f"{where}: {w}")
    return out


def has_unread_content(reports: list[ExtractionReport]) -> bool:
    """True when some content may exist that the grader could not read."""
    for r in reports:
        if not r.readable:
            return True
        for p in r.pages:
            if p.method == "none" or any("not read" in w or "OCR found little" in w for w in p.warnings):
                return True
    return False
